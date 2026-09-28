from __future__ import annotations

import base64, binascii, json, mimetypes, os, re, shutil, subprocess, tempfile, zipfile
from copy import deepcopy
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path
from lxml import etree as LET

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
GENERATED = ROOT / "generated"
RUNTIME = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"
NODE = Path(os.environ.get("MRMG_NODE", RUNTIME / "node/bin/node"))
NODE_MODULES = Path(os.environ.get("MRMG_NODE_MODULES", RUNTIME / "node/node_modules"))
SOURCE_DOCM = ROOT / "templates" / "source.docm"
BLUE, NAVY, INK, LIGHT_BLUE, YELLOW, MAGENTA = "365C73", "203864", "111111", "D9E2F3", "FFF200", "C000A0"
ATTESTATION_KEYS = ["accurate", "eu", "controls", "risk", "scope"]
BUSINESS_RULES = {
    "Credit and Fraud Risk": {
        "Account Receivable / Billed Business": ("≤ $1B", "$1B–$10B", "> $10B"),
        "Adverse Action Volume": ("≤ 1M", "1M–3M", "> 3M"),
        "Write-off": ("≤ $4M", "$4M–$50M", "> $50M"),
    },
    "Marketing": {
        "Gross Contribution Margin": ("≤ $5M", "$5M–$11M", "> $11M"),
        "Pre-Tax Income": ("≤ $5M", "$5M–$11M", "> $11M"),
        "Customers / Prospects Scored": ("≤ 1M", "1M–3M", "> 3M"),
    },
    "Technology and Servicing": {
        "Pre-Tax Income": ("≤ $5M", "$5M–$11M", "> $11M"),
        "Customers / Prospects Scored": ("≤ 1M", "1M–3M", "> 3M"),
    },
    "Finance and Treasury": {
        "Account Receivable / Billed Business": ("≤ $1B", "$1B–$10B", "> $10B"),
        "Balance Sheet": ("≤ $1B", "$1B–$10B", "> $10B"),
        "Write-off": ("≤ $4M", "$4M–$50M", "> $50M"),
    },
    "Compliance and Financial Crimes": {
        "Alert Volume": ("≤ 1M", "1M–3M", "> 3M"),
        "Compliance Review Volume": ("≤ 1M", "1M–3M", "> 3M"),
    },
    "Other": {
        "Customers / Prospects Scored": ("≤ 1M", "1M–3M", "> 3M"),
        "Pre-Tax Income": ("≤ $5M", "$5M–$11M", "> $11M"),
    },
}


def routing(data: dict) -> dict:
    if data.get("solutionType") == "general":
        return {"tier": "Low", "needs_assessment": False, "reason": "General-purpose solutions route directly to Low impact and Section 1."}
    if data.get("regulatory") == "no" and data.get("purpose") in {"customerExperience", "efficiency"}:
        return {"tier": "Low", "needs_assessment": False, "reason": "This combination routes directly to Low impact and does not require Section 2."}
    return {"tier": "Assessment required", "needs_assessment": True, "reason": "The selected use requires the Section 2 impact assessment."}


def assess(data: dict) -> dict:
    routed = routing(data)
    if not routed["needs_assessment"]:
        return {"tier": "Low", "score": None, "early_exit": True, "components": None, "reason": routed["reason"]}

    qn_val = 3 if data.get("regulatory") == "yes" else {"small": 1, "medium": 2, "large": 3}.get(data.get("impactThreshold"), 0)
    reliance_val = {"direct": 3, "human": 2, "multiple": 1}.get(data.get("reliance"), 0)
    purpose_val = 2 if data.get("purpose") in {"core", "peopleCompliance"} else 1
    qualitative_total = reliance_val + purpose_val
    ql_val = 3 if qualitative_total == 5 else 2 if qualitative_total == 4 else 1
    cx_val = int(data.get("explainable") == "no") + int(data.get("fineTuned") == "yes") + int(data.get("multiCall") == "yes")
    dd_val = {"none": 1, "some": 2, "many": 3}.get(data.get("downstream"), 0)
    mic = round(0.4 * qn_val + 0.4 * ql_val + 0.1 * cx_val + 0.1 * dd_val, 2)
    tier = tier_for_mic(mic)
    return {
        "tier": tier, "score": mic, "early_exit": False,
        "components": {"qn": qn_val, "ql": ql_val, "cx": cx_val, "dd": dd_val},
        "formula": "0.4 × qn + 0.4 × ql + 0.1 × cx + 0.1 × dd",
        "reason": "MIC calculated with the approved impact rules.",
    }


def tier_for_mic(mic: float) -> str:
    return "Low" if mic < 1.2 else "Medium" if mic < 1.8 else "High" if mic < 2.2 else "Critical"


def as_list(value) -> list:
    return value if isinstance(value, list) else [value] if value not in (None, "") else []


def prompt_calls(data: dict) -> list[dict]:
    fields = ["promptCallName", "promptCallPurpose", "promptCallText", "promptCallConstraints", "promptCallOutputFormat", "promptCallExampleInput", "promptCallExampleOutput", "promptCallVersion"]
    values = {field: as_list(data.get(field)) for field in fields}
    count = max((len(value) for value in values.values()), default=0)
    return [{field: values[field][index] if index < len(values[field]) else "" for field in fields} for index in range(count)]


def metrics(data: dict) -> list[dict]:
    fields = ["metricName", "metricValue", "metricRationale", "metricGreen", "metricAmber", "metricRed"]
    if data.get("metricName"):
        values = {field: as_list(data.get(field)) for field in fields}
        count = max((len(value) for value in values.values()), default=0)
        return [{field: values[field][index] if index < len(values[field]) else "" for field in fields} for index in range(count)]
    result = []
    for number in (1, 2):
        item = {
            "metricName": data.get(f"metric{number}", ""),
            "metricValue": data.get(f"metric{number}Value", ""),
            "metricRationale": data.get(f"metric{number}Rationale", ""),
            "metricGreen": data.get(f"metric{number}Green", ""),
            "metricAmber": data.get(f"metric{number}Amber", ""),
            "metricRed": data.get(f"metric{number}Red", ""),
        }
        if any(str(value).strip() for value in item.values()): result.append(item)
    return result


def validate(data: dict) -> list[str]:
    required = ["solutionType"]
    if data.get("solutionType") != "general":
        required += ["regulatory", "purpose"]
    required += ["useCaseName", "modelOwner", "modelOwnerEmail", "businessUnit", "implementationDate", "overview", "modelInputs", "generatedOutputs", "modelName", "modelVersion", "hosting", "agentic", "sampleSize", "monitoringFrequency", "businessOwnerName", "businessOwnerTitle", "businessOwnerEmail"]
    if routing(data)["needs_assessment"]:
        required += ["endUsers", "businessProcess", "quantDriver", "impactThreshold", "reliance", "explainable", "fineTuned", "multiCall", "downstream"]
    missing = [name for name in required if not str(data.get(name, "")).strip()]
    submitted_metrics = metrics(data)
    if not submitted_metrics or any(not str(value).strip() for metric in submitted_metrics for value in metric.values()):
        missing.append("metrics")
    calls = prompt_calls(data)
    if not calls:
        missing.append("promptCalls")
    elif any(not str(value).strip() for call in calls for value in call.values()):
        missing.append("promptCalls")
    if data.get("multiCall") == "yes" and len(calls) < 3:
        missing.append("promptCalls")
    markets = data.get("markets", [])
    markets = markets if isinstance(markets, list) else [markets] if markets else []
    if not markets: missing.append("markets")
    if "Other" in markets and not str(data.get("otherMarket", "")).strip(): missing.append("otherMarket")
    selected_attestations = data.get("attestation", [])
    selected_attestations = selected_attestations if isinstance(selected_attestations, list) else [selected_attestations]
    if not all(key in selected_attestations for key in ATTESTATION_KEYS): missing.append("attestation")
    if routing(data)["needs_assessment"]:
        process = data.get("businessProcess")
        driver = data.get("quantDriver")
        if process not in BUSINESS_RULES or driver not in BUSINESS_RULES.get(process, {}):
            missing.append("quantDriver")
    try:
        if data.get("sampleSize") and int(data["sampleSize"]) < 1: missing.append("sampleSize")
    except (TypeError, ValueError):
        missing.append("sampleSize")
    if data.get("businessOwnerEmail") and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", str(data["businessOwnerEmail"])):
        missing.append("businessOwnerEmail")
    if data.get("modelOwnerEmail") and not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", str(data["modelOwnerEmail"])):
        missing.append("modelOwnerEmail")
    for item in as_list(data.get("supportingFiles")):
        if not isinstance(item, dict) or not str(item.get("name", "")).strip() or not str(item.get("data", "")).strip():
            missing.append("supportingFiles"); break
        try:
            base64.b64decode(item["data"], validate=True)
        except (binascii.Error, ValueError):
            missing.append("supportingFiles"); break
    return list(dict.fromkeys(missing))


def shade(cell, color):
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), color); cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar"); tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None: node = OxmlElement(f"w:{name}"); tc_mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr(); borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None: borders = OxmlElement("w:tcBorders"); tc_pr.append(borders)
    for edge, spec in edges.items():
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None: tag = OxmlElement(f"w:{edge}"); borders.append(tag)
        tag.set(qn("w:val"), spec.get("val", "single")); tag.set(qn("w:sz"), str(spec.get("sz", 6))); tag.set(qn("w:color"), spec.get("color", "808080"))


def set_paragraph_border(paragraph, color=MAGENTA, size=16, space=7):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr"); p_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), str(size)); left.set(qn("w:space"), str(space)); left.set(qn("w:color"), color)
    borders.append(left)


def cell_text(cell, text, bold=False, color=INK, size=10, italic=False, align=None):
    cell.text = ""; p = cell.paragraphs[0]
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    display_text = "" if text is None else str(text)
    run = p.add_run(display_text); run.bold = bold; run.italic = italic; run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string(color); run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margins(cell)


def base_document(title=None, subtitle=None):
    doc = Document(); sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.65); sec.left_margin = sec.right_margin = Inches(.9)
    normal = doc.styles["Normal"]; normal.font.name = "Arial"; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.0
    for style_name, size, color in [("Title", 23, BLUE), ("Heading 1", 15.5, BLUE), ("Heading 2", 12, INK)]:
        style = doc.styles[style_name]; style.font.name = "Arial"; style.font.size = Pt(size); style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name == "Heading 2"
    if title:
        p = doc.add_paragraph(title, style="Title"); p.paragraph_format.space_after = Pt(5)
    if subtitle:
        p = doc.add_paragraph(subtitle); p.paragraph_format.space_after = Pt(12)
        for r in p.runs: r.font.name = "Arial"; r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string("666666")
    return doc


def heading(doc, title, level=1):
    p = doc.add_heading(title, level=level); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(5); return p


def option_line(doc, selected, text, indent=.33, italic=False):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(indent); p.paragraph_format.first_line_indent = Inches(-.03); p.paragraph_format.space_after = Pt(4)
    if selected: set_paragraph_border(p)
    box = p.add_run("☒" if selected else "☐"); box.font.name = "DejaVu Sans"; box.font.size = Pt(11)
    run = p.add_run("   " + text); run.font.name = "Arial"; run.font.size = Pt(10.5); run.italic = italic
    return p


def prompt(doc, number, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"{number}.   {text}"); r.bold = True; r.font.name = "Arial"; r.font.size = Pt(10.5); return p


def line_value(doc, label, value, label_width=1.6):
    table = doc.add_table(rows=1, cols=2); table.autofit = False
    left, right = table.rows[0].cells; left.width = Inches(label_width); right.width = Inches(5.5)
    cell_text(left, label, bold=False, size=10.5); cell_text(right, value, size=10.5)
    set_cell_border(left, top={"val":"nil"}, bottom={"val":"nil"}, start={"val":"nil"}, end={"val":"nil"})
    set_cell_border(right, bottom={"val":"single", "sz":6, "color":"595959"}, top={"val":"nil"}, start={"val":"single", "sz":14, "color":MAGENTA}, end={"val":"nil"})
    table.rows[0]._tr.get_or_add_trPr(); return table


def grid_table(doc, headers, rows, widths=None, header_fill=NAVY, header_text_color=None):
    table = doc.add_table(rows=1, cols=len(headers)); table.autofit = False
    if widths:
        for cell, width in zip(table.rows[0].cells, widths): cell.width = Inches(width)
    header_text_color = header_text_color or (INK if header_fill == LIGHT_BLUE else "FFFFFF")
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, header_fill)
        cell_text(cell, text, True, header_text_color, 9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_border(cell, top={"color":"7F7F7F"}, bottom={"color":"7F7F7F"}, start={"color":"7F7F7F"}, end={"color":"7F7F7F"})
    header_pr = table.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement("w:tblHeader"); header_pr.append(repeat); header_pr.append(OxmlElement("w:cantSplit"))
    for row in rows:
        cells = table.add_row().cells
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, (cell, value) in enumerate(zip(cells, row)):
            if widths: cell.width = Inches(widths[i])
            cell_text(cell, value, size=9.5)
            set_cell_border(cell, top={"color":"7F7F7F"}, bottom={"color":"7F7F7F"}, start={"color":"7F7F7F"}, end={"color":"7F7F7F"})
    return table


def attachment_card(doc, filename, description, embed_index=None):
    table = doc.add_table(rows=1, cols=2); table.autofit = False
    icon, details = table.rows[0].cells; icon.width = Inches(.62); details.width = Inches(6.1)
    embed_index = embed_index or {"Outcome Analysis.xlsx": 1, "Prompt Submission Template.docx": 2, "B70+ Attestation Template.docx": 3, "Ongoing Monitoring Plan.docx": 4}.get(filename)
    shade(icon, "E7E6E6"); cell_text(icon, f"[[EMBED{embed_index}]]" if embed_index else "FILE", True, BLUE, 8, align=WD_ALIGN_PARAGRAPH.CENTER)
    details.text = ""; p = details.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    run = p.add_run(filename); run.bold = True; run.font.name = "Arial"; run.font.size = Pt(10)
    p = details.add_paragraph(description); p.paragraph_format.space_after = Pt(0)
    for run in p.runs: run.font.name = "Arial"; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor.from_string("666666")
    for cell in (icon, details):
        set_cell_border(cell, top={"color":"BFBFBF"}, bottom={"color":"BFBFBF"}, start={"color":"BFBFBF"}, end={"color":"BFBFBF"})
    return table


def add_footnote_marker(paragraph, number):
    run = paragraph.add_run(f"[[FN{number}]]")
    run.font.superscript = True
    return run


def build_prompt_doc(data, path):
    doc = base_document("LLM Prompt Documentation Template", "A Guide for Model Owners to Structure and Document Prompts for LLM Calls")
    heading(doc, "Model Owner Documentation Guidance")
    guidance = [
        ("Purpose", "Clearly define the objective of each LLM call and why it is used in the workflow."),
        ("Prompt Structure", "Document the prompt wording, expected context, and any placeholders (e.g., {{user_input}})."),
        ("Policy and Constraints", "List relevant company policies, content guidelines, or operational rules the LLM must enforce or check."),
        ("Output Format", "Specify the expected LLM output format to ensure consistency and easy integration."),
        ("Examples", "Provide sample user messages and sample outputs for each prompt to illustrate expected behavior."),
        ("Revision History", "Track changes and updates to the prompt documentation for traceability."),
    ]
    for label, body in guidance:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent = Inches(.35); r = p.add_run(f"{label}: "); r.bold = True; p.add_run(body)
    doc.add_paragraph("Model owners are encouraged to regularly review and update prompt documentation to reflect evolving business requirements, policy changes, and improvements in LLM capabilities. Thorough documentation helps ensure transparency, reproducibility, and responsible AI deployment.")
    for index, call in enumerate(prompt_calls(data), 1):
        doc.add_page_break()
        heading(doc, f"LLM Call {index}: {call['promptCallName']}")
        grid_table(doc, ["Section", "Details"], [
            ["Prompt Name", call["promptCallName"]],
            ["Purpose", call["promptCallPurpose"]],
            ["Prompt Structure", call["promptCallText"]],
            ["Policy and Constraints", call["promptCallConstraints"]],
            ["Output Format", call["promptCallOutputFormat"]],
            ["Example Input", call["promptCallExampleInput"]],
            ["Example Output", call["promptCallExampleOutput"]],
        ], [1.7, 5.1], header_fill="F2F2F2", header_text_color=INK)
        heading(doc, "Revision History", level=2)
        grid_table(doc, ["Version", "Date", "Change", "Owner"], [[call["promptCallVersion"], datetime.now().strftime("%m/%d/%Y"), "Submitted through First Line intake", data.get("modelOwner", "")]], [0.8, 1.1, 3.4, 1.5], header_fill="7F7F7F")
    doc.save(path)


def build_monitoring_doc(data, path):
    doc = base_document()
    p = doc.add_paragraph(); r = p.add_run("Example for Ongoing Monitoring Plan"); r.bold = True; r.underline = True; r.font.name = "Arial"; r.font.size = Pt(14)
    doc.add_paragraph(f"Please specify the {data.get('monitoringFrequency','annual').lower()} ongoing monitoring plan and include details (for e.g. sample size, performance metrics like accuracy / hallucination rate / acceptance rate / user satisfaction, etc.).")
    submitted_metrics = metrics(data)
    p = doc.add_paragraph(); p.add_run("Model performance metric(s):").bold = True
    doc.add_paragraph("\n".join(f"{index}. {metric['metricName']} — {metric['metricValue']}" for index, metric in enumerate(submitted_metrics, 1)), style=None)
    doc.add_paragraph("Note: In case the performance metrics being tracked are not aligned with model outcomes analysis, please provide a rationale for the choice of model performance metric(s).")
    p = doc.add_paragraph(); p.add_run("Rationale for model performance metrics: ").bold = True; p.add_run(" ".join(metric["metricRationale"] for metric in submitted_metrics))
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.add_run("Sample size:").bold = True
    doc.add_paragraph(f"For model performance monitoring, {data.get('sampleSize','')} records from production data will be used.")
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.add_run("Thresholds and Action Plan:").bold = True
    action = {
        "Red": "• Root Cause Analysis (RCA)\n• Formal governance escalation\n• Remediate, redevelop, or temporarily restrict use",
        "Amber": "• Increase review cadence\n• Investigate persistent deterioration\n• Document rationale, controls, and next steps",
        "Green": "• Continue planned monitoring\n• Retain testing evidence and reviewer conclusions",
    }
    rows = []
    for metric_data in submitted_metrics:
        metric = metric_data["metricName"]
        for status in ("Red", "Amber", "Green"):
            rows.append([metric, status, metric_data[f"metric{status}"], action[status]])
    table = grid_table(doc, ["Metric", "Status", "Threshold of the Testing / Metric", "Action Plan"], rows, [1.15, .75, 1.75, 3.15])
    for row in table.rows[1:]:
        status = row.cells[1].text.strip(); shade(row.cells[1], "FF0000" if status == "Red" else "FFC000" if status == "Amber" else "00B050")
        for run in row.cells[1].paragraphs[0].runs: run.font.color.rgb = RGBColor(255,255,255); run.bold = True
    doc.save(path)


def build_attestation_doc(data, path):
    doc = base_document("Reference Template for B70+ Attestation")
    p = doc.add_paragraph(); p.add_run("Subject: ").bold = True; p.add_run(f"B70+ Attestation for GenAI Use Case – {data['useCaseName']}")
    doc.add_paragraph(f"Dear {data['businessOwnerName']},")
    doc.add_paragraph("To proceed with model risk certification, a formal attestation from the business owner (B70+) is required.")
    doc.add_paragraph("This attestation serves as confirmation that the use case information, control environment, testing, residual risk, and ongoing monitoring plan have been reviewed and approved.")
    p = doc.add_paragraph(); p.add_run("Please find attached:").bold = True
    for item in ["Completed model documentation", "Outcome testing results", "Ongoing monitoring plan"]: doc.add_paragraph(item, style="List Bullet")
    heading(doc, "Attestation")
    doc.add_paragraph("As the designated business owner, I confirm the following:")
    attestations = [
        ("Use Within Scope", "The model will be used only within the intended scope described in the submitted documentation."),
        ("Mandatory Control Effectiveness and Implementation", "Mandatory controls have been tested for effectiveness and will be implemented in production."),
        ("Testing Effectiveness", "The testing performed is appropriate for the use case and supports the conclusions documented in the submission."),
        ("Residual Risk Acceptance", "Residual risk is understood, accepted, and within the business risk appetite."),
        ("Ongoing Monitoring", f"Ongoing monitoring will be performed at least annually; the submitted plan currently specifies a {str(data['monitoringFrequency']).lower()} cadence."),
    ]
    for number, (name, item) in enumerate(attestations, 1):
        p = doc.add_paragraph(); p.add_run(f"{number}. {name}\n").bold = True; p.add_run(item)
    p = doc.add_paragraph("Mandatory controls are defined in the framework"); p.runs[0].italic = True; add_footnote_marker(p, 1)
    heading(doc, "Approval", level=2)
    line_value(doc, "Name:", data["businessOwnerName"]); line_value(doc, "Email:", data["businessOwnerEmail"]); line_value(doc, "Title:", data["businessOwnerTitle"]); line_value(doc, "Date:", datetime.now().strftime("%m/%d/%Y"))
    doc.save(path)


def build_b70_email(data, path, attachments):
    message = EmailMessage()
    message["From"] = f"{data['modelOwner']} <{data['modelOwnerEmail']}>"
    message["To"] = f"{data['businessOwnerName']} <{data['businessOwnerEmail']}>"
    message["Subject"] = f"B70+ Attestation for GenAI Use Case - {data['useCaseName']}"
    message["X-Unsent"] = "1"
    lines = [
        f"Dear {data['businessOwnerName']},", "",
        "To proceed with model risk certification, a formal attestation from the business owner (B70+) is required.", "",
        "This attestation serves as confirmation that:",
        "1. The model will be used only within the intended scope described in the documentation.",
        "2. Mandatory controls have been tested for effectiveness and will be implemented in production.",
        "3. Testing performed is appropriate for the use case and supports the documented conclusions.",
        "4. Residual risk is understood, accepted, and within the business risk appetite.",
        f"5. Ongoing monitoring will be completed at least annually; the submitted plan specifies a {str(data['monitoringFrequency']).lower()} cadence.", "",
        "Please find attached:", "- Completed model documentation", "- Outcome testing results", "- Ongoing monitoring plan", "",
        "Please reply confirming your approval of the above attestation.", "",
        "Regards,", data["modelOwner"], data["modelOwnerEmail"], "",
    ]
    message.set_content("\n".join(lines))
    for attachment in attachments:
        content_type = attachment["content_type"]
        maintype, subtype = content_type.split("/", 1)
        message.add_attachment(Path(attachment["path"]).read_bytes(), maintype=maintype, subtype=subtype, filename=attachment["filename"])
    path.write_bytes(message.as_bytes())


def add_submission_content(doc, data, result):
    routed = routing(data)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Who should complete: "); r.bold = True; r.italic = True
    p.add_run("Use-case owner or modeling lead (B40+)").italic = True
    p = doc.add_paragraph(); r = p.add_run("Instructions: "); r.bold = True; r.italic = True
    p.add_run("For each question, choose the single best option. If unsure, choose the higher-impact option.").italic = True

    prompt(doc, 1, "Which of the following types of GenAI solution fits this use case best? (Select one)")
    option_line(doc, data.get("solutionType") == "general", "General-purpose AI capability or tool, e.g., ChatGPT, Microsoft Office Copilot, coding accelerators, or another capability without a specific business purpose.")
    option_line(doc, data.get("solutionType") == "custom", "Customized solution for a specific business purpose or use, e.g., task generation, control recommender, help-desk assistant, etc.")
    if data.get("solutionType") != "general":
        prompt(doc, 2, "Is the model output used as an input to any regulatory reporting (e.g., CCAR, CECL, SAR filings, stress testing, capital or liquidity reporting, etc.)?")
        option_line(doc, data.get("regulatory") == "yes", "Yes"); option_line(doc, data.get("regulatory") == "no", "No")
        prompt(doc, 3, "What is the purpose and business usage of the model?")
        purpose_options = [
            ("core", "Model decision has direct impact on AXP's core business, e.g., payment and lending decisions."),
            ("peopleCompliance", "Model is used in applications that affect employment, compliance, legal, or regulatory decisions."),
            ("customerExperience", "Model enhances customer experience, communications, or service interactions."),
            ("efficiency", "Model creates productivity or efficiency within AXP's core business, e.g., code suggestions or colleague support."),
        ]
        for key, text in purpose_options: option_line(doc, data.get("purpose") == key, text, italic=True)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
    run = p.add_run("Next Step:"); run.bold = True
    shade_run = OxmlElement("w:shd"); shade_run.set(qn("w:fill"), YELLOW); run._r.get_or_add_rPr().append(shade_run)
    message = " It is a Low Impact model. Proceed to fill in Section 1. Submit this document to Model Store after completing Section 1." if not routed["needs_assessment"] else " Model Impact Category assessment required. Complete Section 1 and Section 2, then submit the document on Model Store and contact MRMG for next steps."
    r = p.add_run(message); r.bold = True; r.font.color.rgb = RGBColor.from_string("548235" if not routed["needs_assessment"] else BLUE)

    doc.add_page_break()
    heading(doc, "Section 1: GenAI Low Impact Model Document Questionnaire")
    heading(doc, "Model Details", level=2)
    markets = data.get("markets", []); markets = markets if isinstance(markets, list) else [markets]
    p = doc.add_paragraph(); p.add_run("Select all markets in which the model is currently used or expected to be used at go-live (select all that apply):").bold = True
    option_line(doc, "US" in markets, "US"); option_line(doc, "EU" in markets, "EU*")
    option_line(doc, "Other" in markets, "Other – Please specify: " + (data.get("otherMarket", "") if "Other" in markets else ""))
    p = doc.add_paragraph("Customer-facing routing and impact are evaluated under the framework"); p.runs[0].italic = True; p.runs[0].font.size = Pt(9); add_footnote_marker(p, 1)

    prompt(doc, 1, "Please provide the name and an overview of the use case along with a summary of model inputs and outputs.")
    line_value(doc, "Use Case Name:", data.get("useCaseName")); line_value(doc, "Overview:", data.get("overview"))
    grid_table(doc, ["Model Inputs", "Generated Outputs"], [[data.get("modelInputs"), data.get("generatedOutputs")]], [3.35, 3.35], header_fill=LIGHT_BLUE)
    p = doc.add_paragraph(); p.add_run("Optional Supporting documents").bold = True
    supporting = as_list(data.get("supportingFiles"))
    doc.add_paragraph("If applicable, attach supporting workflow or control documentation. " + ("Included in the ZIP: " + ", ".join(str(item.get("name", "")) for item in supporting) if supporting else "No optional supporting files were supplied."))
    for index, item in enumerate(supporting, 5):
        attachment_card(doc, Path(str(item.get("name", "supporting-file"))).name, "Optional supporting evidence supplied with this submission", embed_index=index)

    prompt(doc, 2, "GenAI technology: Please provide the Generative AI model leveraged, hosting location, and whether the use case is implemented as an Agentic AI.")
    line_value(doc, "Model Name:", data.get("modelName")); line_value(doc, "Model Version:", data.get("modelVersion"))
    hosting = "Model inference runs within an AXP-governed platform" if data.get("hosting") == "axp" else "Third-party governed platform where data handling is governed outside AXP"
    line_value(doc, "Hosting Location:", hosting)
    p = doc.add_paragraph(); p.add_run("Agentic AI?").bold = True
    option_line(doc, data.get("agentic") == "yes", "Yes"); option_line(doc, data.get("agentic") == "no", "No")
    prompt(doc, 3, "Please attach the model validation results, including:")
    for item in ["Prompts (applicable for in-house developed use cases)", "Sample size reviewed", "Performance indicators used (e.g., model accuracy, output quality, user acceptance)", "Outcome analysis"]:
        doc.add_paragraph(item, style="List Bullet")
    attachment_card(doc, "Outcome Analysis.xlsx", "Completed validation results and metric rationale")
    attachment_card(doc, "Prompt Submission Template.docx", "Completed prompt wording, constraints, outputs, and revision history")
    prompt(doc, 4, "What is the planned implementation date for the use case?")
    line_value(doc, "Implementation Date:", data.get("implementationDate"))

    heading(doc, "5. Additional Details Required")
    heading(doc, "B70+ Business Attestation", level=2)
    doc.add_paragraph("Attach an email from the business B70+ confirming that:")
    attestations = [
        "All information provided regarding the use case in the model documentation is accurate.",
        "If used in the EU, the use does not fall under Prohibited AI practices or High-risk AI systems under the EU AI Act.",
        "Mandatory controls have been tested for effectiveness and will be implemented in production.",
        "Residual risk is understood, accepted, and within the business risk appetite.",
        "The model will be used only within the intended scope described in this document.",
    ]
    for item in attestations: doc.add_paragraph(item, style="List Bullet")
    attachment_card(doc, "B70+ Attestation Template.docx", "Completed Word attestation from the named B70+ business owner")
    doc.add_paragraph("The attestation may be provided by the primary B70+ business owner. Where there is no single owner, such as for a foundational capability, a B70+ owner who uses the capability or owns the relevant process may provide the attestation.")
    heading(doc, "Ongoing Monitoring Plan", level=2)
    doc.add_paragraph("Submit the ongoing monitoring plan with defined metrics, thresholds, cadence, action triggers, and evidence-retention expectations.")
    attachment_card(doc, "Ongoing Monitoring Plan.docx", "Completed monitoring cadence, thresholds, and action plan")
    doc.add_page_break(); p = heading(doc, "Mandatory Controls", level=2); add_footnote_marker(p, 2)
    for item in [
        "User access controls and appropriate AI-use disclaimers.",
        "Incident reporting, fallback, and backup options.",
        "Approved upstream models and governed hosting arrangements.",
        "Controls preventing sensitive-data leakage and harmful content.",
        "Robust implementation, testing, monitoring, and change-management controls.",
    ]: doc.add_paragraph(item, style="List Bullet")

    if routed["needs_assessment"]:
        doc.add_page_break(); heading(doc, "Section 2: Model Risk Tier Assessment")
        p = doc.add_paragraph(); p.add_run("Who are the intended end users of this use case? (Select one)").bold = True
        option_line(doc, data.get("endUsers") == "none", "No direct end users / foundational capability")
        option_line(doc, data.get("endUsers") == "customer", "Customer-facing or applied to customer-impacting decisions")
        option_line(doc, data.get("endUsers") == "internal", "Internal colleagues")
        heading(doc, "A. Business Impact (Quantitative)", level=2)
        impact_rows = []
        for process, drivers in BUSINESS_RULES.items():
            for driver, bands in drivers.items():
                chosen = process == data.get("businessProcess") and driver == data.get("quantDriver")
                for band_name, threshold, key in (("Large", bands[2], "large"), ("Medium", bands[1], "medium"), ("Small", bands[0], "small")):
                    impact_rows.append([f"{'☒' if chosen else '☐'} {process}", f"{'☒' if chosen else '☐'} {driver}", f"{'☒' if chosen and data.get('impactThreshold') == key else '☐'} {band_name}: {threshold}"])
        grid_table(doc, ["Business Process", "Quantitative Driver (annual)", "Quantitative Threshold"], impact_rows, [1.95, 2.25, 2.6], header_fill=LIGHT_BLUE)
        p = doc.add_paragraph("All quantitative driver metrics—including action volume, gross contribution margin, pre-tax income, and related measures—are measured on an annual basis."); p.runs[0].italic = True; p.runs[0].font.size = Pt(9)
        p = doc.add_paragraph("Adverse Action Volume is the number of customers adversely impacted. Alert Volume is measured by entities screened for screening models and by alerts or cases evaluated for true-match/false-positive models. If no traditional metric is available, estimate and report pre-tax income impact."); p.runs[0].italic = True; p.runs[0].font.size = Pt(9)
        heading(doc, "B. Business Importance (Qualitative)", level=2)
        direct_p = option_line(doc, data.get("reliance") == "direct", "Direct reliance / automated decision"); add_footnote_marker(direct_p, 3)
        option_line(doc, data.get("reliance") == "human", "Human review or fallback")
        option_line(doc, data.get("reliance") == "multiple", "Multiple reviews / recommendation only")
        heading(doc, "C. Model Complexity", level=2)
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.add_run("C1 – Explainability feasible?").bold = True
        first = option_line(doc, data.get("explainable") == "yes", "Yes"); first.paragraph_format.keep_with_next = True; option_line(doc, data.get("explainable") == "no", "No")
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.add_run("C2 – Foundational model fine-tuned?").bold = True
        first = option_line(doc, data.get("fineTuned") == "yes", "Yes"); first.paragraph_format.keep_with_next = True; option_line(doc, data.get("fineTuned") == "no", "No")
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.add_run("C3 – More than two sequential LLM calls?").bold = True
        first = option_line(doc, data.get("multiCall") == "yes", "Yes"); first.paragraph_format.keep_with_next = True; option_line(doc, data.get("multiCall") == "no", "No")
        heading(doc, "D. Interdependency", level=2)
        option_line(doc, data.get("downstream") == "none", "0–1 downstream dependencies")
        option_line(doc, data.get("downstream") == "some", "2–5 downstream dependencies")
        option_line(doc, data.get("downstream") == "many", "6+ downstream dependencies")
        components = result["components"]
        grid_table(doc, ["qn (40%)", "ql (40%)", "cx (10%)", "dd (10%)", "MIC"], [[components["qn"], components["ql"], components["cx"], components["dd"], f"{result['score']:.2f}"]], [1.25, 1.25, 1.25, 1.25, 1.3], header_fill=LIGHT_BLUE)
        p = doc.add_paragraph(); set_paragraph_border(p); p.add_run("Model Impact Category: ").bold = True; p.add_run(result["tier"]).bold = True
        p = doc.add_paragraph("MIC = 0.4 × qn + 0.4 × ql + 0.1 × cx + 0.1 × dd. Bands: Low < 1.2; Medium < 1.8; High < 2.2; Critical ≥ 2.2.")
        p.runs[0].font.size = Pt(9); p.runs[0].italic = True

    doc.add_page_break(); heading(doc, "Appendix: Sample Pre-Tax Income Estimation")
    doc.add_paragraph("For productivity and efficiency use cases, estimate annual pre-tax income impact using documented time saved, affected colleague or customer volumes, adoption, and applicable loaded cost or value assumptions. Retain the calculation and assumptions with the submission.")

    heading(doc, "Supporting Files in Submission ZIP", level=2)
    doc.add_paragraph("These completed artifacts are included as separate, usable files in the submission ZIP:")
    for name in ["Outcome Analysis.xlsx", "Prompt Submission Template.docx", "Ongoing Monitoring Plan.docx", "B70+ Attestation Template.docx", "B70+ Attestation Email.eml", "submission.json"]:
        p = doc.add_paragraph(style="List Bullet"); p.add_run(name).bold = True


FOOTNOTE_TEXT = {
    1: "Customer-facing GenAI use cases are assessed based on whether the model output is used for direct decisioning in AXP’s core lending and payments business, as such models may lead to potential customer impact or harm. The extent of impact is determined through the structured risk assessment in Section 2, while non-core business use cases that do not affect customers’ ability to access credit or make payments are considered to have minimal customer impact or harm.",
    2: "Mandatory controls include but are not limited to: User access control; disclaimers to inform use of AI; incident reporting or backup options in case of discontinuation or disruption of service; usage of approved upstream models (if applicable); prevention of sensitive data leakage and blocking of harmful content, e.g., AI Firewall; and robust implementation and change management controls, including segregated development, test, and production environments, release approvals, and a rollback path.",
    3: "Model outputs that directly inform or result in a business decision or action must be subject to review by a subject-matter expert for each applicable case prior to use. For outputs that are informational, advisory, or otherwise do not drive business decisions or actions, subject-matter expert review may be performed on a sample to ascertain model accuracy.",
}


def patch_word_package(path, embeddings=(), footnote_texts=None):
    footnote_texts = FOOTNOTE_TEXT if footnote_texts is None else footnote_texts
    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "v": "urn:schemas-microsoft-com:vml",
        "o": "urn:schemas-microsoft-com:office:office",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
        "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    }
    with tempfile.TemporaryDirectory() as unpack:
        root = Path(unpack)
        with zipfile.ZipFile(path) as zf: zf.extractall(root)
        document_path = root / "word/document.xml"
        document = LET.parse(str(document_path)); document_root = document.getroot()
        for number in footnote_texts:
            for text_node in document_root.xpath(f"//w:t[text()='[[FN{number}]]']", namespaces=ns):
                run = text_node.getparent()
                for child in list(run):
                    if child.tag != f"{{{ns['w']}}}rPr": run.remove(child)
                ref = LET.SubElement(run, f"{{{ns['w']}}}footnoteReference"); ref.set(f"{{{ns['w']}}}id", str(number))
        for index, item in enumerate(embeddings, 1):
            marker = f"[[EMBED{index}]]"
            for text_node in document_root.xpath(f"//w:t[text()='{marker}']", namespaces=ns):
                run = text_node.getparent()
                for child in list(run):
                    if child.tag != f"{{{ns['w']}}}rPr": run.remove(child)
                ext = item["filename"].rsplit(".", 1)[-1].lower()
                prog_id = "Excel.Sheet.12" if ext == "xlsx" else "Word.Document.12" if ext == "docx" else "Package"
                shape_id = f"_x0000_i{1100 + index}"
                object_xml = LET.fromstring(f'''<w:object xmlns:w="{ns['w']}" xmlns:r="{ns['r']}" xmlns:v="{ns['v']}" xmlns:o="{ns['o']}" w:dxaOrig="1200" w:dyaOrig="650"><v:shape id="{shape_id}" type="#_x0000_t75" style="width:60pt;height:34pt" o:ole=""><v:fill color="E7E6E6"/><v:stroke color="7F7F7F"/><v:textbox inset="2pt,2pt,2pt,2pt"><w:txbxContent><w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:rPr><w:b/><w:color w:val="{BLUE}"/><w:sz w:val="16"/></w:rPr><w:t>{ext.upper()}</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape><o:OLEObject Type="Embed" ProgID="{prog_id}" ShapeID="{shape_id}" DrawAspect="Icon" ObjectID="_{1200000000 + index}" r:id="rIdEmbed{index}"/></w:object>''')
                run.append(object_xml)
        document.write(str(document_path), encoding="UTF-8", xml_declaration=True, standalone=True)

        footnotes = LET.Element(f"{{{ns['w']}}}footnotes", nsmap={"w": ns["w"]})
        for note_id, note_type in ((-1, "separator"), (0, "continuationSeparator")):
            node = LET.SubElement(footnotes, f"{{{ns['w']}}}footnote"); node.set(f"{{{ns['w']}}}id", str(note_id)); node.set(f"{{{ns['w']}}}type", note_type)
            p = LET.SubElement(node, f"{{{ns['w']}}}p"); run = LET.SubElement(p, f"{{{ns['w']}}}r"); LET.SubElement(run, f"{{{ns['w']}}}{note_type}")
        for note_id, note_text in footnote_texts.items():
            note = LET.SubElement(footnotes, f"{{{ns['w']}}}footnote"); note.set(f"{{{ns['w']}}}id", str(note_id))
            p = LET.SubElement(note, f"{{{ns['w']}}}p"); run = LET.SubElement(p, f"{{{ns['w']}}}r"); LET.SubElement(run, f"{{{ns['w']}}}footnoteReference")
            text_run = LET.SubElement(p, f"{{{ns['w']}}}r"); text = LET.SubElement(text_run, f"{{{ns['w']}}}t"); text.text = " " + note_text
        LET.ElementTree(footnotes).write(str(root / "word/footnotes.xml"), encoding="UTF-8", xml_declaration=True, standalone=True)

        rels_path = root / "word/_rels/document.xml.rels"; rels = LET.parse(str(rels_path)); rels_root = rels.getroot()
        rel = LET.SubElement(rels_root, f"{{{ns['pr']}}}Relationship"); rel.set("Id", "rIdFootnotes"); rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"); rel.set("Target", "footnotes.xml")
        embeddings_dir = root / "word/embeddings"; embeddings_dir.mkdir(exist_ok=True)
        content_types = LET.parse(str(root / "[Content_Types].xml")); types_root = content_types.getroot()
        override = LET.SubElement(types_root, f"{{{ns['ct']}}}Override"); override.set("PartName", "/word/footnotes.xml"); override.set("ContentType", "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml")
        for index, item in enumerate(embeddings, 1):
            ext = item["filename"].rsplit(".", 1)[-1].lower(); target_name = f"embedded{index}.{ext}"
            shutil.copy2(item["path"], embeddings_dir / target_name)
            rel = LET.SubElement(rels_root, f"{{{ns['pr']}}}Relationship"); rel.set("Id", f"rIdEmbed{index}"); rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"); rel.set("Target", f"embeddings/{target_name}")
            override = LET.SubElement(types_root, f"{{{ns['ct']}}}Override"); override.set("PartName", f"/word/embeddings/{target_name}"); override.set("ContentType", item["content_type"])
        rels.write(str(rels_path), encoding="UTF-8", xml_declaration=True, standalone=True)
        content_types.write(str(root / "[Content_Types].xml"), encoding="UTF-8", xml_declaration=True, standalone=True)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
            for file in sorted(root.rglob("*")):
                if file.is_file(): zf.write(file, file.relative_to(root))


def build_primary(data, result, path, embeddings):
    if SOURCE_DOCM.exists():
        shutil.copy2(SOURCE_DOCM, path)
        with tempfile.TemporaryDirectory() as tmp:
            appendix = Path(tmp) / "appendix.docx"; doc = base_document("GenAI Model Risk Tiering Framework"); add_submission_content(doc, data, result); doc.save(appendix)
            with tempfile.TemporaryDirectory() as unpack:
                root = Path(unpack)
                with zipfile.ZipFile(path) as zf: zf.extractall(root)
                with zipfile.ZipFile(appendix) as zf: app_xml = LET.fromstring(zf.read("word/document.xml"))
                source_path = root / "word/document.xml"; source_xml = LET.parse(str(source_path)); source_root = source_xml.getroot(); w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"; source_body = source_root.find(f"{{{w}}}body"); app_body = app_xml.find(f"{{{w}}}body"); sect = source_body.find(f"{{{w}}}sectPr"); pos = list(source_body).index(sect) if sect is not None else len(source_body)
                for element in list(app_body):
                    if element.tag != f"{{{w}}}sectPr": source_body.insert(pos, deepcopy(element)); pos += 1
                source_xml.write(str(source_path), encoding="UTF-8", xml_declaration=True, standalone=True)
                with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
                    for item in sorted(root.rglob("*")):
                        if item.is_file(): zf.write(item, item.relative_to(root))
    else:
        doc = base_document("GenAI Model Risk Tiering Framework"); add_submission_content(doc, data, result); doc.save(path)
    patch_word_package(path, embeddings)


def generate(data: dict, output_dir: Path | None = None) -> dict:
    missing = validate(data)
    if missing: raise ValueError("Missing required fields: " + ", ".join(missing))
    result = assess(data); enriched = {**data, "metrics": metrics(data), "impactTier": result["tier"], "assessmentScore": result["score"], "assessmentComponents": result["components"], "section2Included": routing(data)["needs_assessment"]}
    slug = re.sub(r"[^a-z0-9]+", "-", data["useCaseName"].lower()).strip("-")[:48] or "submission"
    output_dir = output_dir or GENERATED / f"{datetime.now():%Y%m%d-%H%M%S}-{slug}"; output_dir.mkdir(parents=True, exist_ok=True)
    prompt, monitoring, workbook, attestation_path, b70_email = output_dir / "Prompt Submission Template.docx", output_dir / "Ongoing Monitoring Plan.docx", output_dir / "Outcome Analysis.xlsx", output_dir / "B70+ Attestation Template.docx", output_dir / "B70+ Attestation Email.eml"
    primary = output_dir / f"MRMG First Line Submission{'.docm' if SOURCE_DOCM.exists() else '.docx'}"; json_path = output_dir / "submission.json"
    serializable = {key: value for key, value in enriched.items() if key != "supportingFiles"}
    serializable["supportingFileNames"] = [item.get("name") for item in as_list(enriched.get("supportingFiles"))]
    json_path.write_text(json.dumps(serializable, indent=2), encoding="utf-8"); build_prompt_doc(enriched, prompt); build_monitoring_doc(enriched, monitoring); build_attestation_doc(enriched, attestation_path); patch_word_package(attestation_path, footnote_texts={1: FOOTNOTE_TEXT[2]})
    env = os.environ.copy(); env["NODE_PATH"] = str(NODE_MODULES)
    subprocess.run([str(NODE), str(ROOT / "generate_workbook.mjs"), str(json_path), str(workbook)], check=True, env=env, capture_output=True, text=True)
    optional_paths = []
    if as_list(enriched.get("supportingFiles")):
        supporting_dir = output_dir / "Supporting Documents"; supporting_dir.mkdir(exist_ok=True)
        for index, item in enumerate(as_list(enriched.get("supportingFiles")), 1):
            safe_name = Path(str(item["name"])).name or f"supporting-file-{index}"
            target = supporting_dir / safe_name
            if target.exists(): target = supporting_dir / f"{index}-{safe_name}"
            target.write_bytes(base64.b64decode(item["data"], validate=True)); optional_paths.append(target)
    embeddings = [
        {"filename": workbook.name, "path": workbook, "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        {"filename": prompt.name, "path": prompt, "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
        {"filename": attestation_path.name, "path": attestation_path, "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
        {"filename": monitoring.name, "path": monitoring, "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    ]
    embeddings.extend({"filename": item.name, "path": item, "content_type": mimetypes.guess_type(item.name)[0] or "application/octet-stream"} for item in optional_paths)
    build_primary(enriched, result, primary, embeddings)
    build_b70_email(enriched, b70_email, [
        {"filename": primary.name, "path": primary, "content_type": "application/vnd.ms-word.document.macroEnabled.12" if primary.suffix == ".docm" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
        {"filename": workbook.name, "path": workbook, "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
        {"filename": monitoring.name, "path": monitoring, "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    ])
    package = output_dir / f"{slug}-mrmg-submission.zip"
    with zipfile.ZipFile(package, "w", zipfile.ZIP_DEFLATED) as zf:
        for item in [primary, workbook, prompt, monitoring, attestation_path, b70_email, json_path]: zf.write(item, item.name)
        for item in optional_paths: zf.write(item, f"Supporting Documents/{item.name}")
    return {"directory": output_dir, "package": package, "primary": primary, "workbook": workbook, "prompt": prompt, "monitoring": monitoring, "attestation": attestation_path, "email": b70_email, "result": result}


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(); parser.add_argument("input_json", type=Path); parser.add_argument("output_dir", type=Path); args = parser.parse_args()
    payload = json.loads(args.input_json.read_text(encoding="utf-8")); generated = generate(payload, args.output_dir)
    print(json.dumps({k: str(v) if isinstance(v, Path) else v for k, v in generated.items()}, indent=2))
