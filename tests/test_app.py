import json
from email import policy
from email.parser import BytesParser
import tempfile
import threading
import unittest
import urllib.request
import zipfile
from http.server import ThreadingHTTPServer
from pathlib import Path
from xml.etree import ElementTree as ET
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import generator
from generator import BUSINESS_RULES, assess, generate, metrics, routing, tier_for_mic, validate
from server import Handler
from docx import Document


def sample(profile="low"):
    data = {
        "solutionType": "custom", "regulatory": "no", "purpose": "efficiency",
        "useCaseName": "Knowledge Assistant", "modelOwner": "Model Team", "businessUnit": "Enterprise Services",
        "implementationDate": "2026-09-30", "markets": ["US"], "overview": "Answers questions from approved knowledge.",
        "modelInputs": "User query and approved documents", "generatedOutputs": "Grounded answer with citations",
        "modelName": "OpenAI GPT", "modelVersion": "5.1", "hosting": "axp", "agentic": "no",
        "sampleSize": "1447", "metric1": "Accuracy", "metric1Value": "82%",
        "metric1Rationale": "Each query maps to one expected answer.", "metric2": "Hallucination rate", "metric2Value": "3%",
        "metric2Rationale": "Unsupported claims are material to this workflow.",
        "promptCallName": "Retrieval answer", "promptCallPurpose": "Answer from approved context",
        "promptCallText": "Answer only from the supplied context and cite the source.",
        "promptCallConstraints": "Do not use outside knowledge.", "promptCallOutputFormat": "Grounded prose with citations",
        "promptCallExampleInput": "What is the policy?", "promptCallExampleOutput": "The policy states… [source]", "promptCallVersion": "1.0",
        "metric1Green": "Accuracy >= 80%", "metric1Amber": "65% <= Accuracy < 80%", "metric1Red": "Accuracy < 65%",
        "metric2Green": "Hallucination rate <= 3%", "metric2Amber": "3% < Hallucination rate <= 5%", "metric2Red": "Hallucination rate > 5%",
        "monitoringFrequency": "Annual", "attestation": ["accurate", "eu", "controls", "risk", "scope"],
        "businessOwnerName": "Avery Morgan", "businessOwnerTitle": "SVP, Enterprise Services", "businessOwnerEmail": "avery.morgan@example.com",
    }
    profiles = {
        "low_assessed": {"purpose": "core", "endUsers": "none", "businessProcess": "Technology and Servicing", "quantDriver": "Customers / Prospects Scored", "impactThreshold": "small", "reliance": "multiple", "explainable": "yes", "fineTuned": "no", "multiCall": "no", "downstream": "none"},
        "medium": {"purpose": "core", "endUsers": "internal", "businessProcess": "Technology and Servicing", "quantDriver": "Customers / Prospects Scored", "impactThreshold": "small", "reliance": "human", "explainable": "no", "fineTuned": "yes", "multiCall": "no", "downstream": "none"},
        "high": {"purpose": "core", "endUsers": "customer", "businessProcess": "Credit and Fraud Risk", "quantDriver": "Adverse Action Volume", "impactThreshold": "medium", "reliance": "direct", "explainable": "yes", "fineTuned": "no", "multiCall": "no", "downstream": "none"},
        "critical": {"regulatory": "yes", "purpose": "core", "endUsers": "customer", "businessProcess": "Credit and Fraud Risk", "quantDriver": "Adverse Action Volume", "impactThreshold": "large", "reliance": "direct", "explainable": "no", "fineTuned": "yes", "multiCall": "yes", "downstream": "many"},
    }
    if profile in profiles:
        data.update(profiles[profile])
    if profile == "critical":
        for field in ["promptCallName", "promptCallPurpose", "promptCallText", "promptCallConstraints", "promptCallOutputFormat", "promptCallExampleInput", "promptCallExampleOutput", "promptCallVersion"]:
            data[field] = [data[field], data[field] + " 2", data[field] + " 3"]
    return data


class RulesTests(unittest.TestCase):
    def test_vba_early_exit_paths(self):
        self.assertFalse(routing(sample("low"))["needs_assessment"])
        general = sample("low"); general["solutionType"] = "general"; general.pop("regulatory"); general.pop("purpose")
        self.assertFalse(routing(general)["needs_assessment"])
        self.assertTrue(assess(general)["early_exit"])
        self.assertIsNone(assess(general)["score"])

    def test_exact_mic_profiles_and_bands(self):
        expected = {"low_assessed": (0.9, "Low"), "medium": (1.5, "Medium"), "high": (2.1, "High"), "critical": (3.0, "Critical")}
        for profile, (mic, tier) in expected.items():
            with self.subTest(profile=profile):
                result = assess(sample(profile)); self.assertEqual(result["score"], mic); self.assertEqual(result["tier"], tier)
        self.assertEqual([(value, tier_for_mic(value)) for value in (1.19, 1.2, 1.79, 1.8, 2.19, 2.2)], [(1.19, "Low"), (1.2, "Medium"), (1.79, "Medium"), (1.8, "High"), (2.19, "High"), (2.2, "Critical")])

    def test_regulatory_reporting_forces_quantitative_three(self):
        data = sample("low_assessed"); data["regulatory"] = "yes"
        self.assertEqual(assess(data)["components"]["qn"], 3)

    def test_required_fields_and_conditional_other_market(self):
        data = sample("low"); data.pop("modelName"); data["markets"] = ["US", "Other"]
        missing = validate(data)
        self.assertIn("modelName", missing); self.assertIn("otherMarket", missing)
        data = sample("low"); data["attestation"] = ["accurate"]
        self.assertIn("attestation", validate(data))

    def test_business_process_driver_compatibility(self):
        data = sample("high"); data["quantDriver"] = "Alert Volume"
        self.assertIn("quantDriver", validate(data))
        self.assertEqual(BUSINESS_RULES["Credit and Fraud Risk"]["Adverse Action Volume"], ("≤ 1M", "1M–3M", "> 3M"))

    def test_multicall_requires_three_documented_calls(self):
        data = sample("high"); data["multiCall"] = "yes"
        self.assertIn("promptCalls", validate(data))

    def test_one_metric_is_valid_and_metrics_are_repeatable(self):
        one = sample("low")
        for field in ["metric2", "metric2Value", "metric2Rationale", "metric2Green", "metric2Amber", "metric2Red"]:
            one.pop(field)
        self.assertNotIn("metrics", validate(one))
        three = sample("low")
        three.update({
            "metricName": ["Accuracy", "Hallucination rate", "User satisfaction"],
            "metricValue": ["82%", "3%", "4.5/5"],
            "metricRationale": ["Measures correctness", "Measures unsupported claims", "Measures utility"],
            "metricGreen": [">= 80%", "<= 3%", ">= 4/5"],
            "metricAmber": ["65-79%", "3-5%", "3-3.9/5"],
            "metricRed": ["< 65%", "> 5%", "< 3/5"],
        })
        self.assertNotIn("metrics", validate(three))
        self.assertEqual([item["metricName"] for item in metrics(three)], ["Accuracy", "Hallucination rate", "User satisfaction"])


class GenerationTests(unittest.TestCase):
    def test_complete_package_documents_rules_and_section2(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = generate(sample("high"), Path(tmp))
            for key in ["package", "primary", "workbook", "prompt", "monitoring", "attestation", "email"]:
                self.assertTrue(result[key].exists(), key); self.assertGreater(result[key].stat().st_size, 300)
            with zipfile.ZipFile(result["primary"]) as zf:
                document_xml = zf.read("word/document.xml").decode("utf-8")
                self.assertIn("Section 2: Model Impact Category Assessment", document_xml)
                self.assertIn("C000A0", document_xml)
                self.assertIn("1M–3M", document_xml)
                self.assertIn("Business Impact Threshold Reference", document_xml)
            primary_doc = Document(result["primary"])
            component_tables = [table for table in primary_doc.tables if table.rows[0].cells[0].text == "qn (40%)"]
            self.assertEqual(len(component_tables), 1)
            self.assertEqual(component_tables[0].rows[1].cells[2].text, "0")
            with zipfile.ZipFile(result["workbook"]) as zf:
                self.assertIn("xl/workbook.xml", zf.namelist())
                sheet_xml = b"".join(zf.read(name) for name in zf.namelist() if name.startswith("xl/worksheets/sheet"))
                for label in [b"Sample Size", b"Performance Metric 1", b"Rationale for Metric 1", b"Performance Metric 2"]:
                    self.assertIn(label, sheet_xml)
            attestation_text = "\n".join(paragraph.text for paragraph in Document(result["attestation"]).paragraphs)
            self.assertIn("Mandatory controls have been tested", attestation_text)
            self.assertIn("at least annually", attestation_text)
            email_message = BytesParser(policy=policy.default).parsebytes(result["email"].read_bytes())
            self.assertEqual(email_message["From"], "Avery Morgan <avery.morgan@example.com>")
            self.assertIn("B70+ Attestation", email_message["Subject"])
            self.assertIn("Residual risk is understood", email_message.get_content())
            with zipfile.ZipFile(result["package"]) as zf:
                self.assertEqual(len(zf.namelist()), 7)
                self.assertIn("B70+ Attestation Template.docx", zf.namelist())
                self.assertIn("B70+ Attestation Email.eml", zf.namelist())

    def test_low_document_omits_section2(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = generate(sample("low"), Path(tmp))
            with zipfile.ZipFile(result["primary"]) as zf:
                document_xml = zf.read("word/document.xml").decode("utf-8")
            self.assertNotIn("Section 2: Model Impact Category Assessment", document_xml)
            self.assertIn("5. Additional Details Required", document_xml)

    def test_source_docm_vba_part_is_preserved(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp); source = tmp_path / "source.docm"; Document().save(source)
            with zipfile.ZipFile(source, "a", zipfile.ZIP_DEFLATED) as zf: zf.writestr("word/vbaProject.bin", b"synthetic-vba-marker")
            previous = generator.SOURCE_DOCM; generator.SOURCE_DOCM = source
            try: result = generator.generate(sample("low"), tmp_path / "out")
            finally: generator.SOURCE_DOCM = previous
            self.assertEqual(result["primary"].suffix, ".docm")
            with zipfile.ZipFile(result["primary"]) as zf: self.assertEqual(zf.read("word/vbaProject.bin"), b"synthetic-vba-marker")


class HttpTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True); cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown(); cls.server.server_close(); cls.thread.join(timeout=2)

    def test_health_and_generation(self):
        base = f"http://127.0.0.1:{self.server.server_port}"
        with urllib.request.urlopen(base + "/api/health") as response: payload = json.loads(response.read())
        self.assertEqual(payload["status"], "ok")
        request = urllib.request.Request(base + "/api/generate", data=json.dumps(sample("medium")).encode(), headers={"Content-Type": "application/json"}, method="POST")
        with urllib.request.urlopen(request, timeout=30) as response: generated = json.loads(response.read())
        self.assertIn("MIC 1.50", generated["summary"])
        with urllib.request.urlopen(base + generated["download"], timeout=30) as response: package = response.read()
        self.assertGreater(len(package), 1000)


if __name__ == "__main__":
    unittest.main()
